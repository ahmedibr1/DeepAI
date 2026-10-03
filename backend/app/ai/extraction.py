"""Text extraction from opportunity documents, keeping the provenance every citation needs.

Native text only: each extracted block records where it came from (page, slide, sheet, section), so a
finding can say "RFP.pdf — Page 37 — Section 4.2" instead of an unsupported claim. Scanned pages have no
text layer; they are reported as needing OCR rather than guessed at (see `ExtractionResult.ocr_pages`).
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from typing import BinaryIO

HEADING = re.compile(r"^\s*((?:\d+\.){0,4}\d+)\s+(\S.{0,120})$")     # "4.2 Integration requirements"
MIN_CHARS_PER_PAGE = 40   # below this a PDF page is treated as scanned


@dataclass
class Block:
    """One piece of text with the location it came from."""
    text: str
    page_number: int | None = None
    slide_number: int | None = None
    sheet_name: str | None = None
    section_path: str | None = None
    kind: str = "text"     # text | table | notes | heading


@dataclass
class ExtractionResult:
    blocks: list[Block] = field(default_factory=list)
    page_count: int | None = None
    ocr_pages: list[int] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def text_length(self) -> int:
        return sum(len(b.text) for b in self.blocks)


def _track_section(line: str, current: list[str]) -> str | None:
    m = HEADING.match(line)
    if not m or len(line) > 160:
        return None
    number, title = m.group(1), m.group(2).strip()
    depth = number.count(".") + 1
    del current[depth - 1:]
    current.append(f"{number} {title}")
    return " › ".join(current)


def extract_pdf(stream: BinaryIO) -> ExtractionResult:
    from pypdf import PdfReader

    reader = PdfReader(stream)
    result = ExtractionResult(page_count=len(reader.pages))
    section: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:                      # a damaged page must not stop the document
            result.warnings.append(f"Page {index} could not be read ({exc.__class__.__name__}).")
            continue
        if len(text.strip()) < MIN_CHARS_PER_PAGE:
            result.ocr_pages.append(index)
            continue
        current_section = None
        for line in text.splitlines():
            found = _track_section(line, section)
            current_section = found or current_section or (" › ".join(section) or None)
        result.blocks.append(Block(text=text.strip(), page_number=index, section_path=current_section))
    return result


def extract_docx(stream: BinaryIO) -> ExtractionResult:
    from docx import Document

    doc = Document(stream)
    result = ExtractionResult()
    section: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            result.blocks.append(Block(text="\n".join(buffer).strip(), section_path=" › ".join(section) or None))
            buffer.clear()

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower()
        if style.startswith("heading"):
            flush()
            depth = int(style.replace("heading", "").strip() or 1)
            del section[depth - 1:]
            section.append(text)
            continue
        _track_section(text, section)
        buffer.append(text)
    flush()

    for number, table in enumerate(doc.tables, start=1):
        rows = ["\t".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if any(r.strip() for r in rows):
            result.blocks.append(Block(text="\n".join(rows), section_path=f"Table {number}", kind="table"))
    return result


def extract_pptx(stream: BinaryIO) -> ExtractionResult:
    from pptx import Presentation

    prs = Presentation(stream)
    result = ExtractionResult(page_count=len(prs.slides))
    for index, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        title = None
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                text = shape.text_frame.text.strip()
                if title is None and shape == slide.shapes.title:
                    title = text
                parts.append(text)
            if getattr(shape, "has_table", False):
                parts.extend("\t".join(c.text.strip() for c in row.cells) for row in shape.table.rows)
        if parts:
            result.blocks.append(Block(text="\n".join(parts), slide_number=index, section_path=title))
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        if notes:
            result.blocks.append(Block(text=notes, slide_number=index, section_path=title, kind="notes"))
    return result


def extract_xlsx(stream: BinaryIO, max_rows: int = 2000) -> ExtractionResult:
    from openpyxl import load_workbook

    wb = load_workbook(stream, read_only=True, data_only=True)
    result = ExtractionResult()
    for sheet in wb.worksheets:
        rows: list[str] = []
        for count, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if count > max_rows:
                result.warnings.append(f"Sheet “{sheet.title}” was read up to row {max_rows}.")
                break
            cells = ["" if c is None else str(c).strip() for c in row]
            if any(cells):
                rows.append("\t".join(cells))
        if rows:
            result.blocks.append(Block(text="\n".join(rows), sheet_name=sheet.title, kind="table"))
    wb.close()
    return result


def extract_text(stream: BinaryIO) -> ExtractionResult:
    raw = stream.read()
    for encoding in ("utf-8", "utf-8-sig", "cp1256", "latin-1"):   # cp1256 covers Windows Arabic files
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        return ExtractionResult(warnings=["The file could not be decoded as text."])
    section: list[str] = []
    blocks: list[Block] = []
    current: list[str] = []
    current_path: str | None = None

    def flush() -> None:
        if current:
            blocks.append(Block(text="\n".join(current).strip(), section_path=current_path))
            current.clear()

    for line in text.splitlines():
        heading = _track_section(line, section)
        if heading:                      # a numbered heading starts a new block, so citations stay precise
            flush()
            current_path = heading
        if line.strip():
            current.append(line.rstrip())
    flush()
    return ExtractionResult(blocks=blocks)


def extract_csv(stream: BinaryIO) -> ExtractionResult:
    raw = stream.read().decode("utf-8", errors="replace")
    rows = ["\t".join(cell.strip() for cell in row) for row in csv.reader(io.StringIO(raw)) if any(row)]
    return ExtractionResult(blocks=[Block(text="\n".join(rows), kind="table")] if rows else [])


EXTRACTORS = {
    "pdf": extract_pdf, "docx": extract_docx, "pptx": extract_pptx, "xlsx": extract_xlsx,
    "txt": extract_text, "eml": extract_text, "csv": extract_csv,
}
UNSUPPORTED_NOTE = {
    "doc": "Legacy .doc is not read directly. Save it as .docx and upload again.",
    "xls": "Legacy .xls is not read directly. Save it as .xlsx and upload again.",
    "ppt": "Legacy .ppt is not read directly. Save it as .pptx and upload again.",
    "msg": "Outlook .msg is not read directly. Export the email as PDF and upload again.",
    "png": "Images have no text layer; OCR is not enabled in this deployment.",
    "jpg": "Images have no text layer; OCR is not enabled in this deployment.",
    "jpeg": "Images have no text layer; OCR is not enabled in this deployment.",
}


def extract(extension: str, stream: BinaryIO) -> ExtractionResult:
    extractor = EXTRACTORS.get(extension.lower())
    if extractor is None:
        return ExtractionResult(warnings=[UNSUPPORTED_NOTE.get(extension.lower(), f".{extension} is not supported.")])
    try:
        return extractor(stream)
    except Exception as exc:
        return ExtractionResult(warnings=[f"The file could not be read ({exc.__class__.__name__})."])
